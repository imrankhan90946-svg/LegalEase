"""Export-format regression tests."""
from io import BytesIO
from pathlib import Path

from docx import Document

from document_utils.docx_generator import generate_docx
from document_utils.pdf_generator import generate_pdf
from document_utils.txt_generator import generate_txt

SAMPLE = """FREELANCE SERVICES AGREEMENT

PARTIES
Jane Doe (Provider) and TechNova Inc. (Client).

TERMS AND CONDITIONS
Payment within 30 days; Confidentiality must be maintained; Project delivery by agreed deadline.

SIGNATURES
Provider: ____________________    Client: ____________________"""


def test_txt_is_utf8_and_normalized():
    payload = generate_txt("Agreement — ‘draft’…")
    assert payload.decode("utf-8") == 'Agreement - \'draft\'...\n'


def test_docx_is_valid_editable_word_document_and_splits_terms():
    payload = generate_docx(SAMPLE, "Freelance Contract")
    document = Document(BytesIO(payload))
    text = "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert "FREELANCE SERVICES AGREEMENT" in text
    assert "Payment within 30 days" in text
    assert any(paragraph.style.name == "List Bullet" for paragraph in document.paragraphs)


def test_pdf_is_valid_pdf():
    payload = generate_pdf(SAMPLE, "Freelance Contract")
    assert payload.startswith(b"%PDF-")
    assert len(payload) > 500


def test_uploaded_logo_embeds_in_both_exports():
    logo = (Path(__file__).resolve().parents[1] / "assets" / "logo.png").read_bytes()
    docx_payload = generate_docx(SAMPLE, "Freelance Contract", logo=logo)
    pdf_payload = generate_pdf(SAMPLE, "Freelance Contract", logo=logo)
    document = Document(BytesIO(docx_payload))
    assert any("image" in relationship.reltype for relationship in document.sections[0].header.part.rels.values())
    assert pdf_payload.startswith(b"%PDF-")


def test_empty_exports_fail_usefully():
    import pytest
    with pytest.raises(ValueError, match="empty"):
        generate_docx("  ", "Contract")
    with pytest.raises(ValueError, match="empty"):
        generate_pdf("  ", "Contract")
